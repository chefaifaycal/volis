from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.section import WD_ORIENT
d=Document('raw.docx')
NAVY=RGBColor(0x1F,0x3A,0x5F); TEAL=RGBColor(0x2A,0x7F,0x8E)
for sec in d.sections:
    sec.page_width=Cm(21.59); sec.page_height=Cm(27.94)
    for m in ('left_margin','right_margin'): setattr(sec,m,Cm(2))
    sec.top_margin=Cm(1.8); sec.bottom_margin=Cm(1.8)
    p=sec.footer.paragraphs[0]; p.text='Volis — Cahier des charges · Volet technique et sécurité · 11 octobre 2026'
    for r in p.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x80,0x80,0x80)
def setfont(style,name='Calibri',size=None,color=None,bold=None):
    st=d.styles[style]; st.font.name=name
    rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
    if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
    for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),name)
    for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme','w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if size: st.font.size=Pt(size)
    if color: st.font.color.rgb=color
    if bold is not None: st.font.bold=bold
for sn in ['Normal','Body Text','First Paragraph','Compact']:
    try: setfont(sn,size=10.5)
    except KeyError: pass
for sn,sz,c in [('Title',22,NAVY),('Subtitle',12,TEAL),('Heading 1',16,NAVY),('Heading 2',13,TEAL),('Heading 3',11.5,TEAL)]:
    try: setfont(sn,size=sz,color=c,bold=(sn!='Subtitle'))
    except KeyError: pass
for sn in ('Source Code','Verbatim Char'):
    try: setfont(sn,name='Consolas',size=8.5)
    except KeyError: pass
def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd')
    sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc); tcPr.append(sh)
def borders(tbl):
    tblPr=tbl._tbl.tblPr; b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'B8C7D0'); b.append(el)
    tblPr.append(b)
from docx.shared import Emu
TOTAL=Cm(17.5)
def widths(t):
    n=len(t.columns); L=[0]*n
    for row in t.rows:
        for i,c in enumerate(row.cells[:n]):
            L[i]+=min(len(c.text),400)
    L=[max(x,1)**0.7 for x in L]
    L=[max(x/sum(L),0.6/n) for x in L]; S=sum(L)
    return [int(TOTAL*x/S) for x in L]
for t in d.tables:
    ws=widths(t)
    tblPr=t._tbl.tblPr
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    for e in tblPr.findall(qn('w:tblW')): tblPr.remove(e)
    tw=OxmlElement('w:tblW'); tw.set(qn('w:w'),str(int(TOTAL/635))); tw.set(qn('w:type'),'dxa'); tblPr.append(tw)
    grid=t._tbl.tblGrid
    for gc,w in zip(grid.findall(qn('w:gridCol')),ws): gc.set(qn('w:w'),str(int(w/635)))
    for row in t.rows:
        for c,w in zip(row.cells,ws): c.width=Emu(w)
    borders(t)
    for ri,row in enumerate(t.rows):
        for c in row.cells:
            if ri==0: shade(c,'1F3A5F')
            elif ri%2==0: shade(c,'EAF2F5')
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1); p.paragraph_format.space_before=Pt(1)
                for r in p.runs:
                    r.font.size=Pt(8.5)
                    if ri==0: r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); r.font.bold=True
    # repeat header
    trPr=t.rows[0]._tr.get_or_add_trPr(); h=OxmlElement('w:tblHeader'); h.set(qn('w:val'),'true'); trPr.append(h)
# block quotes shading
for p in d.paragraphs:
    if p.style.name in ('Block Text',):
        pPr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:fill'),'FFF4E0'); pPr.append(sh)
d.core_properties.title='Volis — Cahier des charges, section : Volet technique et sécurité'
d.core_properties.author='Faycal Chefai'
d.core_properties.language='fr-CA'
d.save('/workspace/volis/cdc-technique-securite.docx')
