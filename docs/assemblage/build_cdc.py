# Assemblage du « Cahier des charges Volis » : relit les 3 sections à chaque exécution.
import re, os, subprocess, json, shutil, hashlib, datetime
import pypandoc
W = '/workspace/volis'; A = f'{W}/assemblage'
SRC = [('A', 'Spécifications fonctionnelles et UX', 'Équipe UI/UX', '/workspace/volis-design/specs-fonctionnelles-ux.md'),
       ('B', 'Volet technique et sécurité', 'Équipe technique', f'{W}/cdc-technique-securite.md'),
       ('C', 'Volet légal de la monétisation', 'Équipe juridique', f'{W}/cahier-legal-monetisation.md')]
DATE = '11 octobre 2026'; VERSION = '1.0'
TEAL, INK, CREAM, CORAL = '#0F766E', '#1E2A3A', '#F5F1E8', '#FF7A59'

def strip_front(s):
    if s.startswith('---'):
        e = s.index('\n---', 3); return s[e+4:].lstrip('\n')
    return s
def demote(s, n=1):
    out, code = [], False
    for line in s.split('\n'):
        if line.startswith('```'): code = not code
        if not code and re.match(r'^#{1,5} ', line): line = '#'*n + line
        out.append(line)
    return '\n'.join(out)
def mermaid_to_img(s, part, fmt):
    blocks = re.findall(r'```mermaid\n(.*?)```', s, re.S)
    caps = ['Schéma du système', 'Machine d’états d’une livraison']
    for i, b in enumerate(blocks, 1):
        png = f'{A}/diagramme-{part}{i}.png'
        old = f'{W}/cdc/d{i}.mmd'
        if os.path.exists(old) and open(old).read().strip() == b.strip() and os.path.exists(f'{W}/cdc/d{i}.png'):
            shutil.copy(f'{W}/cdc/d{i}.png', png)
        else:
            open(f'{A}/d{i}.mmd', 'w').write(b)
            subprocess.run(['npx', '-y', '@mermaid-js/mermaid-cli', '-i', f'{A}/d{i}.mmd', '-o', png, '-s', '3', '-b', 'white', '-p', f'{W}/cdc/pp.json'], capture_output=True, timeout=300)
        cap = caps[i-1] if i <= len(caps) else f'Diagramme {i}'
        w = '16.5cm' if i == 1 else '9.5cm'
        s = s.replace(f'```mermaid\n{b}```', f'![Figure {part}-{i} — {cap}]({png}){{width={w}}}', 1)
    return s
HARMONISE = [  # renvois entre sections, une fois assemblées (édition mineure, sens inchangé)
    ("Ce détail correspond à l'écran de confirmation prévu par Ui Designer.", "Ce détail correspond à l'écran de paiement de la partie A (écran 11)."),
    ("Il doit être assemblé avec la section de l'équipe UI (fonctionnalités et parcours) et avec la section de l'équipe juridique (`cahier-legal-monetisation.md`, « Volet légal de la monétisation »).",
     "Il est assemblé avec la partie A (spécifications fonctionnelles et UX) et la partie C (volet légal de la monétisation) du présent cahier des charges."),
]
def load():
    parts = []
    for code, title, team, path in SRC:
        raw = open(path, encoding='utf-8').read()
        meta = re.search(r'^subtitle:\s*"(.*?)"', raw, re.M)
        s = strip_front(raw)
        for a, b in HARMONISE: s = s.replace(a, b)
        s = mermaid_to_img(s, code, 'img')
        s = demote(s, 1)
        parts.append(dict(code=code, title=title, team=team, sub=meta.group(1) if meta else '', md=s, path=path,
                          mtime=datetime.datetime.fromtimestamp(os.path.getmtime(path)).strftime('%Y-%m-%d %H:%M'), sha=hashlib.sha1(raw.encode()).hexdigest()[:8]))
    return parts
INTRO = f'''# Introduction

## Objet du document

Ce cahier des charges décrit l’application **Volis** (nom de travail) : une plateforme de mise en relation entre des **expéditeurs** qui veulent faire livrer un colis à l’étranger et des **voyageurs** vérifiés qui ont de la place dans leurs bagages sur un vol. Le paiement est retenu (séquestre) jusqu’à la remise au destinataire. Le document sert de référence commune pour la conception, le développement, l’estimation et la validation juridique de la première version de l’application (iOS, Android et web).

## Contexte

- **Porteur du projet :** Faycal Chefai (Trois-Rivières, Québec).
- **Marché de départ :** diaspora installée au Québec qui envoie des colis vers sa famille (premiers corridors illustrés : Montréal YUL → Alger ALG, Tlemcen, Oran, Casablanca). Le service s’adresse aux particuliers (consommateurs au sens de la LPC).
- **Promesse :** un envoi plus rapide et moins cher que les transporteurs, avec trois garanties visibles : identité vérifiée (Stripe Identity), paiement retenu jusqu’à la livraison (Stripe Connect) et preuves de remise horodatées.
- **Plateformes :** application mobile Expo / React Native (iOS, Android) et web; backend Supabase hébergé au Canada (région Canada (Central), `ca-central-1`), selon la partie B.
- **Direction visuelle retenue :** concept 1 « Carnet de voyage », avec le logo « Colis caché » et l’icône d’appli « Signature » (charte graphique fournie à part).

## Structure du document

| Partie | Contenu | Préparée par |
|---|---|---|
| **A. Spécifications fonctionnelles et UX** | Rôles, navigation, parcours et histoires utilisateur avec critères d’acceptation, inventaire des 50 écrans, états et cas limites, accessibilité, localisation, événements analytiques, exigences UX de monétisation. | Équipe UI/UX |
| **B. Volet technique et sécurité** | Stack, architecture, modèle de données, machine d’états, webhooks Stripe, exigences de sécurité testables (SEC-01 à SEC-25), durées de conservation, monétisation technique, exigences non fonctionnelles, jalons et risques. | Équipe technique |
| **C. Volet légal de la monétisation** | Sources de revenus, prix et transparence (LPC, Loi sur la concurrence), commission, assurance, versements, séquestre et litiges, marketing et données, clauses des conditions d’utilisation, points à faire valider. | Équipe juridique |

Chaque partie garde sa propre numérotation. Les renvois prennent la forme « partie B, 4.1 » ou, dans le texte des équipes, « volet technique 4.1 » et « volet légal 2 ».

## Conventions de lecture

- **« À confirmer »** : point qui n’a pas été vérifié dans une source officielle ou qui dépend d’un fournisseur (Stripe, Supabase, assureur).
- **« À décider (par Faycal) »** : choix d’affaires ouvert : montants, taux, délais, plafonds. Les valeurs visibles dans les maquettes sont **illustratives**.
- **« À faire valider par l’avocat / le comptable »** : point juridique ou fiscal qui doit être validé avant le lancement.
- Les noms de personnes, numéros de vol, codes et montants des exemples sont **fictifs**.
- Les identifiants **US-xx** (histoires utilisateur, partie A) et **SEC-xx** (exigences de sécurité, partie B) servent à la traçabilité dans les tickets et les tests.

## Documents associés

- Maquettes : `volis-ecrans-concept1.pdf` (50 écrans) et planches par parcours.
- Charte graphique : `volis-charte-graphique.pdf` (17 pages).
- Ressources développeurs : `volis-dev-assets-partie1…3.zip` (tokens, thème React Native, icônes, illustrations, logos, icônes d’appli, polices, inventaire des écrans).
'''
DISCLAIMER = '''**Avertissement.** Ce cahier des charges est un document de travail préparé pour Faycal Chefai. Les passages juridiques, fiscaux et financiers sont une **synthèse informative** : ils ne constituent ni un avis juridique, ni un avis fiscal, ni un avis financier. Ils doivent être validés par un avocat membre du Barreau du Québec, et par un comptable pour la fiscalité, avant le lancement. Les informations sur des services tiers (Stripe, Supabase, Expo, Apple, Google) reflètent leur documentation à la date du document et peuvent changer. Les montants, taux et délais cités en exemple ne sont pas des engagements.'''

def master_md(parts, for_docx=False):
    pb = '\n\n'
    md = INTRO
    for p in parts:
        md += pb + f'# Partie {p["code"]} — {p["title"]}\n\n' + f'*{p["sub"] or "Préparé par " + p["team"]}*\n\n' + p['md'].strip() + '\n'
    return md

# ------------------------------------------------------------------ PDF
WM = open('/workspace/logos/final/01-logo-principal/svg/volis-logo-couleur-fond-creme.svg').read()
WM = re.sub(r'<\?xml[^>]*\?>', '', WM)
ILL = '/workspace/volis-design/dev-assets/illustrations/voyageur-valise-colis@2x.png'
CSS = f'''
@page {{ size: letter; margin: 20mm 19mm 20mm 19mm;
  @bottom-left {{ content: "Volis — Cahier des charges · v{VERSION} · {DATE}"; font: 500 8pt Inter; color: #6F6A60; }}
  @bottom-right {{ content: counter(page) " / " counter(pages); font: 600 8pt Inter; color: #6F6A60; }}
}}
@page cover {{ margin: 0; @bottom-left {{ content: none }} @bottom-right {{ content: none }} }}
html {{ -webkit-print-color-adjust: exact; print-color-adjust: exact; }}
body {{ font: 400 9.6pt/1.5 Inter, sans-serif; color: {INK}; margin: 0; }}
.cover {{ page: cover; width: 215.9mm; height: 279.4mm; position: relative; overflow: hidden; background: {CREAM}; page-break-after: always; }}
h1, h2, h3, h4 {{ font-family: 'Plus Jakarta Sans', sans-serif; color: {INK}; break-after: avoid; page-break-after: avoid; }}
h1 {{ font-size: 20pt; line-height: 1.2; margin: 0 0 10pt; padding-bottom: 6pt; border-bottom: 2.5pt solid {TEAL}; break-before: page; page-break-before: always; }}
h1.part {{ string-set: part content(text); }}
h2 {{ font-size: 13.5pt; margin: 16pt 0 6pt; color: {INK}; }}
h3 {{ font-size: 11pt; margin: 12pt 0 4pt; color: {TEAL}; }}
h4 {{ font-size: 10pt; margin: 10pt 0 4pt; color: {TEAL}; }}
p {{ margin: 0 0 6pt; orphans: 3; widows: 3; }}
a {{ color: {TEAL}; text-decoration: none; word-break: break-word; }}
ul, ol {{ margin: 0 0 6pt; padding-left: 16pt; }} li {{ margin: 0 0 2.5pt; }}
blockquote {{ margin: 6pt 0 10pt; padding: 8pt 12pt; background: #EEF6F4; border-left: 3.5pt solid {TEAL}; border-radius: 0 6pt 6pt 0; }}
blockquote p:last-child {{ margin-bottom: 0; }}
code {{ font: 500 8.4pt 'JetBrains Mono', monospace; background: #EFE8DA; padding: 0 2.5pt; border-radius: 3pt; word-break: break-word; }}
pre {{ background: #F7F3EA; border: 1pt solid #E4DACA; border-radius: 6pt; padding: 8pt 10pt; font-size: 7.8pt; line-height: 1.4; white-space: pre-wrap; break-inside: avoid; }}
pre code {{ background: none; padding: 0; font-size: 7.8pt; }}
table {{ width: 100%; border-collapse: collapse; margin: 6pt 0 10pt; font-size: 8.2pt; line-height: 1.4; }}
thead {{ display: table-header-group; }}
th {{ background: {TEAL}; color: #fff; text-align: left; font-weight: 700; padding: 4pt 6pt; }}
td {{ padding: 4pt 6pt; border-bottom: 0.75pt solid #DCD3C2; vertical-align: top; }}
tr {{ break-inside: avoid; }}
tbody tr:nth-child(even) td {{ background: #FBF8F2; }}
figure {{ margin: 8pt 0 12pt; text-align: center; break-inside: avoid; }}
figure img {{ max-width: 100%; max-height: 190mm; }}
figcaption {{ font: 500 8pt Inter; color: #4E5A69; margin-top: 4pt; }}
input[type=checkbox] {{ margin-right: 4pt; }}
ul.task-list, ul:has(> li > input) {{ list-style: none; padding-left: 4pt; }}
.mk {{ position: absolute; font-size: 1px; color: transparent; }}
.toc h1 {{ break-before: page; }}
.toc ol {{ list-style: none; padding: 0; margin: 0; }}
.toc li {{ display: flex; align-items: baseline; gap: 6pt; margin: 0; padding: 2.4pt 0; }}
.toc li.l1 {{ font: 700 10.5pt 'Plus Jakarta Sans'; margin-top: 8pt; color: {INK}; }}
.toc li.l2 {{ font: 400 9pt Inter; padding-left: 14pt; color: #33404F; }}
.toc li .t {{ flex: none; max-width: 82%; }}
.toc li .d {{ flex: 1; border-bottom: 1pt dotted #B9B09E; transform: translateY(-3pt); }}
.toc li .n {{ flex: none; font-variant-numeric: tabular-nums; min-width: 16pt; text-align: right; }}
.partcover {{ break-before: page; page-break-before: always; }}
.meta td:first-child {{ font-weight: 700; width: 32%; }}
'''
def html_fragment(md):
    return pypandoc.convert_text(md, 'html5', format='markdown+pipe_tables+task_lists+implicit_figures-smart', extra_args=['--wrap=none', '--section-divs=false'])
def build_html(parts, pages=None, markers=True):
    body = html_fragment(master_md(parts))
    heads = []
    def sub(m):
        lvl, attrs, text = m.group(1), m.group(2), m.group(3)
        i = len(heads); plain = re.sub('<[^>]+>', '', text)
        heads.append((int(lvl), plain, i))
        cls = ' class="part"' if lvl == '1' and plain.startswith('Partie') else ''
        mk = f'<span class="mk">@@H{i:03d}@@</span>' if markers else ''
        return f'<h{lvl}{attrs}{cls}>{mk}{text}</h{lvl}>'
    body = re.sub(r'<h([12])([^>]*)>(.*?)</h\1>', sub, body, flags=re.S)
    toc = '<section class="toc"><h1 style="break-before:page">Table des matières</h1><ol>' + ''.join(
        f'<li class="l{l}"><span class="t">{t}</span><span class="d"></span><span class="n">{(pages or {}).get(i, "00")}</span></li>' for l, t, i in heads) + '</ol></section>'
    src_rows = ''.join(f'<tr><td>Partie {p["code"]} — {p["title"]}</td><td>{p["team"]}</td><td><code>{os.path.basename(p["path"])}</code><br><span style="color:#6F6A60">version du {p["mtime"]} · empreinte {p["sha"]}</span></td></tr>' for p in parts)
    cover = (f'<section class="cover"><div style="position:absolute;right:-60mm;top:-40mm;width:190mm;height:190mm;border-radius:46% 54% 50% 50%;background:#BFE3DB;opacity:.65"></div>'
             f'<div style="position:absolute;right:-30mm;top:120mm;width:120mm;height:110mm;border-radius:52% 48% 44% 56%;background:#FFC4B2;opacity:.6"></div>'
             f'<img src="file://{ILL}" style="position:absolute;right:22mm;top:52mm;height:150mm"/>'
             f'<div style="position:absolute;left:22mm;top:30mm;width:62mm">{WM.replace("<svg", "<svg style=\'width:100%;height:auto\'", 1)}</div>'
             f'<div style="position:absolute;left:22mm;top:92mm;width:105mm"><div style="font:700 10pt JetBrains Mono;letter-spacing:2pt;color:{TEAL}">CAHIER DES CHARGES</div>'
             f'<div style="font:800 34pt/1.08 \'Plus Jakarta Sans\';margin-top:6pt;letter-spacing:-.5pt">Application Volis</div>'
             f'<div style="font:400 12.5pt/1.45 Inter;color:#4E5A69;margin-top:10pt">Mise en relation d’expéditeurs de colis et de voyageurs vérifiés, avec paiement retenu jusqu’à la livraison. iOS, Android et web.</div></div>'
             f'<div style="position:absolute;left:22mm;bottom:28mm;right:22mm;border-top:2pt solid {INK};padding-top:8pt;display:flex;justify-content:space-between;font:500 9.5pt/1.5 Inter">'
             f'<div style="flex:1.3"><b>Préparé pour</b><br>Faycal Chefai<br>Trois-Rivières (Québec)</div><div style="flex:1"><b>Version</b><br>{VERSION}<br>{DATE}</div><div style="flex:1.4"><b>Contenu</b><br>A. Fonctionnel et UX · B. Technique<br>et sécurité · C. Volet légal</div></div>'
             f'<div style="position:absolute;left:22mm;bottom:14mm;font:500 8pt Inter;color:#6F6A60">Document de travail · ne constitue pas un avis juridique · noms et montants des exemples fictifs</div></section>')
    info = (f'<section><h2 style="margin-top:0">Avertissement</h2>' + html_fragment(DISCLAIMER) +
            f'<h2>Informations sur le document</h2><table class="meta"><tbody><tr><td>Titre</td><td>Cahier des charges — Application Volis</td></tr><tr><td>Version</td><td>{VERSION} ({DATE})</td></tr>'
            f'<tr><td>Préparé pour</td><td>Faycal Chefai, porteur du projet</td></tr><tr><td>Langue</td><td>Français (Canada)</td></tr><tr><td>Statut</td><td>Document de travail, à valider (juridique, fiscal, fournisseurs)</td></tr></tbody></table>'
            f'<h2>Sections assemblées</h2><table><thead><tr><th>Partie</th><th>Préparée par</th><th>Fichier source</th></tr></thead><tbody>{src_rows}</tbody></table>'
            f'<p style="color:#4E5A69;font-size:8.5pt">Les trois sections sont reprises intégralement. Seuls les renvois entre sections ont été harmonisés lors de l’assemblage.</p></section>')
    body = re.sub(r'<colgroup>.*?</colgroup>', '', body, flags=re.S)
    body = re.sub(r'<td>((?:SEC|US)-\d+)</td>', r'<td style="white-space:nowrap">\1</td>', body)
    return f'<!doctype html><html lang="fr-CA"><head><meta charset="utf-8"><title>Volis — Cahier des charges</title><style>{CSS}</style></head><body>{cover}{info}{toc}{body}</body></html>', heads
def chrome_pdf(html_path, pdf_path):
    subprocess.run(['google-chrome', '--headless', '--no-sandbox', '--disable-gpu', '--no-pdf-header-footer', '--virtual-time-budget=5000', f'--print-to-pdf={pdf_path}', f'file://{html_path}'], capture_output=True, timeout=180)
def find_pages(pdf, n):
    txt = subprocess.run(['pdftotext', pdf, '-'], capture_output=True, text=True).stdout.split('\f')
    pages = {}
    for pi, t in enumerate(txt, 1):
        for m in re.finditer(r'@@H(\d{3})@@', t): pages.setdefault(int(m.group(1)), pi)
    return pages
def build_pdf(parts):
    h, heads = build_html(parts); hp = f'{A}/cdc.html'; open(hp, 'w').write(h)
    chrome_pdf(hp, f'{A}/pass1.pdf'); pages = find_pages(f'{A}/pass1.pdf', len(heads))
    missing = [i for _, _, i in heads if i not in pages]
    for it in range(3):
        h, heads = build_html(parts, pages); open(hp, 'w').write(h); chrome_pdf(hp, f'{A}/pass2.pdf')
        p2 = find_pages(f'{A}/pass2.pdf', len(heads))
        if p2 == pages: break
        pages = p2
    # version finale sans marqueurs (positionnés en absolu : la mise en page ne bouge pas)
    h, _ = build_html(parts, pages, markers=False); open(hp, 'w').write(h)
    out = f'{W}/cahier-des-charges-volis.pdf'; chrome_pdf(hp, out)
    return out, pages, missing, heads

# ------------------------------------------------------------------ DOCX
def build_docx(parts):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor, Emu
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    md = master_md(parts, for_docx=True)
    disc = '\n\n::: {custom-style="Avertissement"}\n' + DISCLAIMER + '\n:::\n\n'
    src = ''.join(f'| Partie {p["code"]} — {p["title"]} | {p["team"]} | `{os.path.basename(p["path"])}` (version du {p["mtime"]}) |\n' for p in parts)
    info = ('\n\n# Informations sur le document {.unnumbered}\n\n| | |\n|---|---|\n| Titre | Cahier des charges — Application Volis |\n| Version | ' + VERSION + ' (' + DATE + ') |\n| Préparé pour | Faycal Chefai, porteur du projet |\n| Langue | Français (Canada) |\n| Statut | Document de travail, à valider (juridique, fiscal, fournisseurs) |\n\n'
            '**Sections assemblées**\n\n| Partie | Préparée par | Fichier source |\n|---|---|---|\n' + src + '\nLes trois sections sont reprises intégralement; seuls les renvois entre sections ont été harmonisés lors de l’assemblage.\n')
    full = disc + info + '\n\n' + md
    ref = f'{A}/reference.docx'
    if not os.path.exists(ref):
        subprocess.run([pypandoc.get_pandoc_path(), '-o', ref, '--print-default-data-file', 'reference.docx'], check=True)
    raw = f'{A}/raw.docx'
    pypandoc.convert_text(full, 'docx', format='markdown+pipe_tables+task_lists+fenced_divs+raw_attribute-smart', outputfile=raw,
                          extra_args=['--toc', '--toc-depth=2', f'--metadata=title:Cahier des charges — Application Volis', f'--metadata=subtitle:Version {VERSION} · {DATE} · préparé pour Faycal Chefai',
                                      '--metadata=lang:fr-CA', '--metadata=toc-title:Table des matières', f'--resource-path={A}', f'--reference-doc={ref}'])
    d = Document(raw)
    TEALc, INKc, GREY = RGBColor(0x0F, 0x76, 0x6E), RGBColor(0x1E, 0x2A, 0x3A), RGBColor(0x6F, 0x6A, 0x60)
    def setfont(style, name='Calibri', size=None, color=None, bold=None):
        try: st = d.styles[style]
        except KeyError: return
        st.font.name = name; rpr = st.element.get_or_add_rPr(); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), name)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        if size: st.font.size = Pt(size)
        if color is not None: st.font.color.rgb = color
        if bold is not None: st.font.bold = bold
    for sn in ['Normal', 'Body Text', 'First Paragraph', 'Compact', 'Block Text', 'Table', 'Avertissement']: setfont(sn, size=10)
    for sn, sz, c in [('Title', 26, INKc), ('Subtitle', 12, TEALc), ('Heading 1', 17, INKc), ('Heading 2', 13, TEALc), ('Heading 3', 11.5, TEALc), ('Heading 4', 10.5, TEALc), ('TOC Heading', 17, INKc)]:
        setfont(sn, size=sz, color=c, bold=(sn not in ('Subtitle',)))
    for sn in ('Source Code', 'Verbatim Char'): setfont(sn, name='Consolas', size=8.5)
    # Heading 1 = nouvelle page
    for sn in ('Heading 1',):
        pPr = d.styles[sn].element.get_or_add_pPr(); e = OxmlElement('w:pageBreakBefore'); pPr.append(e)
    for sec in d.sections:
        sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
        sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.9)
        sec.different_first_page_header_footer = True
        p = sec.footer.paragraphs[0]; p.text = f'Volis — Cahier des charges · v{VERSION} · {DATE}\tPage '
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.5), 2)
        r = p.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin'); r._r.append(f1)
        r2 = p.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = ' PAGE '; r2._r.append(it)
        r3 = p.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'end'); r3._r.append(f2)
        for rr in p.runs: rr.font.size = Pt(8); rr.font.color.rgb = GREY; rr.font.name = 'Calibri'
    # tableaux
    def shade(cell, hexc):
        tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexc); tcPr.append(sh)
    for t in d.tables:
        tblPr = t._tbl.tblPr; b = OxmlElement('w:tblBorders')
        for e in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D4CAB7'); b.append(el)
        tblPr.append(b)
        n = len(t.columns); L = [0]*n
        for row in t.rows:
            for i, c in enumerate(row.cells[:n]): L[i] += min(len(c.text), 300)
        L = [max(x, 1)**0.7 for x in L]; L = [max(x/sum(L), 0.5/n) for x in L]; S = sum(L)
        ws = [int(Cm(17.5)*L[i]/S) for i in range(n)]
        for row in t.rows:
            for i, c in enumerate(row.cells[:n]): c.width = Emu(ws[i])
        grid = t._tbl.find(qn('w:tblGrid'))
        if grid is not None:
            for i, gc in enumerate(grid.findall(qn('w:gridCol'))[:n]): gc.set(qn('w:w'), str(int(ws[i]/635)))
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        tw = tblPr.find(qn('w:tblW'))
        if tw is not None: tw.set(qn('w:type'), 'dxa'); tw.set(qn('w:w'), str(int(Cm(17.5)/635)))
        hdr = t.rows[0]
        if any(c.text.strip() for c in hdr.cells):
            for c in hdr.cells:
                shade(c, '0F766E')
                for p in c.paragraphs:
                    for r in p.runs: r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); r.font.bold = True
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    for r in p.runs: r.font.size = Pt(8.5)
    # page de titre : logo + avertissement encadré, saut de page avant la table des matières
    body = d.element.body
    title = next(p for p in d.paragraphs if p.style.name == 'Title')
    lp = title.insert_paragraph_before(); lp.add_run().add_picture('/workspace/volis/assemblage/logo-cover.png', width=Cm(5.5))
    sp = title.insert_paragraph_before(); sp.paragraph_format.space_after = Pt(60)
    kick = title.insert_paragraph_before('CAHIER DES CHARGES'); kick.runs[0].font.color.rgb = TEALc; kick.runs[0].font.bold = True; kick.runs[0].font.size = Pt(10)
    for p in d.paragraphs:
        if p.style.name == 'Avertissement':
            pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr')
            for e in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement(f'w:{e}'); el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '8'); el.set(qn('w:space'), '6'); el.set(qn('w:color'), '0F766E'); bd.append(el)
            pPr.append(bd); sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'EEF6F4'); pPr.append(sh)
            p.paragraph_format.space_before = Pt(140)
            break
    # réordonner : avertissement sur la page de titre, table des matières après « Informations »
    sdt = body.find(qn('w:sdt'))
    av = next(p for p in d.paragraphs if p.style.name == 'Avertissement')
    sdt.addprevious(av._p)
    intro = next(p for p in d.paragraphs if p.style.name == 'Heading 1' and p.text.strip() == 'Introduction')
    intro._p.addprevious(sdt)
    # la TdM commence sur une nouvelle page
    first = sdt.find('.//' + qn('w:p'))
    tocpar = sdt.find(qn('w:sdtContent')).find(qn('w:p'))
    pPr = tocpar.find(qn('w:pPr'))
    if pPr is None: pPr = OxmlElement('w:pPr'); tocpar.insert(0, pPr)
    pPr.append(OxmlElement('w:pageBreakBefore'))
    # Word met la table à jour à l'ouverture
    st = d.settings.element; uf = OxmlElement('w:updateFields'); uf.set(qn('w:val'), 'true'); st.append(uf)
    out = f'{W}/cahier-des-charges-volis.docx'; d.save(out)
    fill_toc(out); return out

def fill_toc(path):
    """Pré-remplit la table des matières (numéros calculés par un rendu LibreOffice) ; Word la recalcule à l'ouverture."""
    from docx import Document
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    from docx.shared import Cm
    tmp = f'{A}/lo'; shutil.rmtree(tmp, ignore_errors=True); os.makedirs(tmp)
    def lo_pages():
        shutil.copy(path, f'{tmp}/x.docx')
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, f'{tmp}/x.docx'], capture_output=True, timeout=240)
        return subprocess.run(['pdftotext', f'{tmp}/x.pdf', '-'], capture_output=True, text=True).stdout.split('\f')
    d = Document(path)
    heads = [(1 if p.style.name == 'Heading 1' else 2, p.text.strip()) for p in d.paragraphs if p.style.name in ('Heading 1', 'Heading 2') and p.text.strip()]
    heads = [h for h in heads if h[1] != 'Informations sur le document']
    def norm(t): return re.sub(r'\s+', ' ', t).strip()
    def compute(pages_txt, start):
        res, cur = [], start
        for lvl, t in heads:
            key = norm(t)[:38]
            for pi in range(cur, len(pages_txt)):
                if key in norm(pages_txt[pi]): cur = pi; break
            res.append(cur + 1)
        return res
    def write(nums):
        d = Document(path)
        sdt = d.element.body.find(qn('w:sdt')); cont = sdt.find(qn('w:sdtContent'))
        ps = cont.findall(qn('w:p')); title_p = ps[0]
        for p in ps[1:]: cont.remove(p)
        def run(text=None, fld=None, instr=None, tab=False):
            r = OxmlElement('w:r')
            if fld: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), fld); r.append(e)
            if instr: e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = instr; r.append(e)
            if tab: r.append(OxmlElement('w:tab'))
            if text is not None: e = OxmlElement('w:t'); e.set(qn('xml:space'), 'preserve'); e.text = text; r.append(e)
            return r
        for k, ((lvl, t), n) in enumerate(zip(heads, nums)):
            p = OxmlElement('w:p'); pPr = OxmlElement('w:pPr'); ps_ = OxmlElement('w:pStyle'); ps_.set(qn('w:val'), f'TOC{lvl}'); pPr.append(ps_)
            tabs = OxmlElement('w:tabs'); tb = OxmlElement('w:tab'); tb.set(qn('w:val'), 'right'); tb.set(qn('w:leader'), 'dot'); tb.set(qn('w:pos'), str(int(17.5/2.54*1440))); tabs.append(tb); pPr.append(tabs)
            if lvl == 2: ind = OxmlElement('w:ind'); ind.set(qn('w:left'), '360'); pPr.append(ind)
            sp = OxmlElement('w:spacing'); sp.set(qn('w:before'), '120' if lvl == 1 else '0'); sp.set(qn('w:after'), '30'); pPr.append(sp)
            p.append(pPr)
            if k == 0:
                p.append(run(fld='begin')); p.append(run(instr=' TOC \\o "1-2" \\h \\z \\u ')); p.append(run(fld='separate'))
            r = run(t);
            if lvl == 1:
                rPr = OxmlElement('w:rPr'); rPr.append(OxmlElement('w:b')); r.insert(0, rPr)
            p.append(r); p.append(run(str(n), tab=True))
            if k == len(heads) - 1: p.append(run(fld='end'))
            cont.append(p)
        d.save(path)
    write(['0'] * len(heads))
    for _ in range(2):
        txt = lo_pages()
        start = next(i for i, t in enumerate(txt) if 'Table des matières' in t) + 1
        nums = compute(txt, start); write(nums)
    return nums


if __name__ == '__main__':
    import sys
    parts = load()
    for p in parts: print(p['code'], p['path'], p['mtime'], p['sha'], len(p['md']))
    open(f'{A}/cahier-des-charges-volis.md', 'w').write('---\ntitle: "Cahier des charges — Application Volis"\nsubtitle: "Version ' + VERSION + ' · ' + DATE + ' · préparé pour Faycal Chefai"\nlang: fr-CA\n---\n\n> ' + DISCLAIMER + '\n\n' + master_md(parts))
    if 'docx' in sys.argv:
        print('DOCX', build_docx(parts))
    if 'pdf' in sys.argv:
        out, pages, missing, heads = build_pdf(parts)
        print('PDF', out, os.path.getsize(out), 'missing', missing, 'last', max(pages.values()))
